"""新版直播经营分析：历史直播表现 × 数据罗盘履约实销。"""
from datetime import date, timedelta
import io
import json

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Font, PatternFill

from core.db import load_product_master
from core.live_analytics import (
    _duration_seconds,
    load_live_actuals,
    load_live_auxiliary,
    load_live_room_data,
    load_live_sessions,
)
from core.theme import page_header
from core.utils import clear_cache_on_page_change
from core.inventory import attach_inventory_summary, render_inventory_detail


DISPLAY_COLUMN_NAMES = {
    "id": "记录ID",
    "live_room_id": "直播场次ID", "shop_name": "直播间／店铺", "anchor_name": "主播",
    "start_time": "开播时间", "duration_seconds": "直播时长（秒）", "style_code": "货号",
    "product_id": "商品ID", "product_name": "商品名称", "click_users": "商品点击人数",
    "sold_units": "成交件数", "paid_amount": "平台支付", "talk_count": "讲解次数",
    "pre_ship_refund_amount": "发货前退款", "post_ship_refund_amount": "发货后退款",
    "module": "指标分组", "metric_name": "指标名称", "metric_value": "指标数值",
    "raw_value": "原始值", "unit": "单位", "benchmark_value": "对标数值",
    "benchmark_raw": "对标原始值", "comparison_display": "对比表现", "channel_name": "流量渠道",
    "avg_watch_duration": "人均观看时长", "watch_count": "观看次数",
    "watch_users": "观看人数", "order_count": "成交订单数", "avg_order_amount": "笔单价",
    "watch_conversion_rate": "观看成交率", "shop_bound_spend": "店铺绑定投放消耗",
    "shop_promoted_spend": "店铺被投投放消耗", "product_image_url": "商品图片",
    "talk_start_epoch": "讲解开始时间戳", "talk_end_epoch": "讲解结束时间戳",
    "talk_duration_seconds": "讲解时长（秒）", "viewer_change": "在线人数变化",
    "avg_online_users": "平均在线人数", "session_start_utc": "开播时间",
    "talk_start_time": "讲解开始时间", "talk_end_time": "讲解结束时间", "ship_amount": "范围发货",
    "return_amount": "范围退货", "net_amount": "范围实销", "sale_date": "销售日期",
    "created_at": "记录时间", "updated_at": "更新时间", "imported_at": "导入时间",
}


def parse_epoch_utc(values: pd.Series) -> pd.Series:
    """同时兼容秒级和毫秒级 Unix 时间戳，统一转换为 UTC。"""
    numeric = pd.to_numeric(values, errors="coerce")
    seconds = pd.to_datetime(numeric.where(numeric.abs() < 100_000_000_000), unit="s", utc=True, errors="coerce")
    milliseconds = pd.to_datetime(numeric.where(numeric.abs() >= 100_000_000_000), unit="ms", utc=True, errors="coerce")
    return seconds.fillna(milliseconds)


def format_china_time(values: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(values, utc=True, errors="coerce")
    return parsed.dt.tz_convert("Asia/Shanghai").dt.strftime("%Y-%m-%d %H:%M:%S").fillna("")


def localize_table(frame: pd.DataFrame) -> pd.DataFrame:
    """只调整页面展示：字段中文化，比例按两位小数显示。"""
    if frame is None:
        return pd.DataFrame()
    display = attach_inventory_summary(frame.copy())
    display = display.drop(
        columns=["id", "live_room_id", "product_id", "imported_at", "created_at", "updated_at"],
        errors="ignore",
    )
    if "talk_start_time" not in display.columns and "talk_start_epoch" in display.columns:
        display["talk_start_time"] = parse_epoch_utc(display["talk_start_epoch"])
    if "talk_end_time" not in display.columns and "talk_end_epoch" in display.columns:
        display["talk_end_time"] = parse_epoch_utc(display["talk_end_epoch"])
    for time_column in ["talk_start_time", "talk_end_time"]:
        if time_column in display.columns:
            display[time_column] = format_china_time(display[time_column])
    display = display.drop(columns=["talk_start_epoch", "talk_end_epoch"], errors="ignore")
    if "讲解分钟" in display.columns:
        display["讲解时长（分钟）"] = pd.to_numeric(display["讲解分钟"], errors="coerce").round(2)
        display = display.drop(columns=["讲解分钟", "talk_duration_seconds"], errors="ignore")
    elif "talk_duration_seconds" in display.columns:
        display["讲解时长（分钟）"] = (
            pd.to_numeric(display["talk_duration_seconds"], errors="coerce") / 60
        ).round(2)
        display = display.drop(columns=["talk_duration_seconds"], errors="ignore")
    style_column = "style_code" if "style_code" in display.columns else "货号" if "货号" in display.columns else None
    image_column_exists = "product_image_url" in display.columns or "商品图片" in display.columns
    image_lookup = globals().get("style_image_lookup", {})
    if style_column and not image_column_exists and image_lookup:
        display["商品图片"] = display[style_column].astype(str).str.strip().str.upper().map(image_lookup)
    display = display.rename(columns=DISPLAY_COLUMN_NAMES)
    if "商品图片" in display.columns:
        display["商品图片"] = display["商品图片"].replace([0, "0", "", "nan", "None"], pd.NA)
    if "货号" in display.columns and "商品图片" in display.columns:
        ordered_columns = list(display.columns)
        ordered_columns.remove("商品图片")
        ordered_columns.insert(ordered_columns.index("货号") + 1, "商品图片")
        display = display[ordered_columns]
    for column in display.columns:
        if "率" not in str(column) and "占比" not in str(column):
            continue
        numeric = pd.to_numeric(display[column], errors="coerce")
        if numeric.notna().any():
            # 保持数值类型，交给表格显示层格式化；否则前端会按百分比文本
            # 的首字符排序，例如错误地把 9% 排在 12% 前面。
            display[column] = numeric
    return display


def show_table(frame: pd.DataFrame) -> None:
    display = localize_table(frame)
    column_config = {}
    if "商品图片" in display.columns:
        column_config["商品图片"] = st.column_config.ImageColumn("商品图片", width="small")
    for column in display.columns:
        if "率" in str(column) or "占比" in str(column):
            if pd.api.types.is_numeric_dtype(display[column]):
                column_config[column] = st.column_config.NumberColumn(column, format="percent")
    st.dataframe(display, width="stretch", hide_index=True, column_config=column_config)


def safe_number(value, default=0.0) -> float:
    """将数据库空值安全转换为数值，避免 pd.NA 参与布尔判断。"""
    try:
        return default if pd.isna(value) else float(value)
    except (TypeError, ValueError):
        return default


def safe_text(value, default="") -> str:
    try:
        return default if pd.isna(value) else str(value)
    except (TypeError, ValueError):
        return default


st.set_page_config(page_title="直播经营分析", layout="wide", initial_sidebar_state="expanded")
clear_cache_on_page_change("live_operations")
page_header("直播经营分析", "历史直播表现 × 数据罗盘履约实销", "LIVE OPERATIONS", "新版")

st.markdown("""
<style>
/* 直播经营分析设计系统：覆盖全站旧样式，统一密度与节奏。 */
section[data-testid="stSidebar"]{display:block!important}
section[data-testid="stSidebar"] [data-testid="stSidebarNav"]{display:none!important}
.main .block-container{max-width:1780px!important;padding:24px 30px 52px!important}
.live-workspace-nav{position:relative;width:100%;box-sizing:border-box;padding:18px 12px 14px;background:linear-gradient(180deg,#061a2e 0%,#08243b 55%,#071c30 100%);border:1px solid rgba(89,211,239,.18);border-radius:12px;box-shadow:0 8px 24px rgba(4,22,39,.10);color:#eaf7fb;overflow:hidden}
.live-workspace-nav::-webkit-scrollbar{width:5px}.live-workspace-nav::-webkit-scrollbar-thumb{background:rgba(77,211,235,.38);border-radius:999px}.live-workspace-nav::-webkit-scrollbar-track{background:transparent}
.live-workspace-nav__brand{padding:4px 10px 20px;border-bottom:1px solid rgba(148,204,220,.15)}
.live-workspace-nav__eyebrow{margin-bottom:6px;color:#49d7ee;font-size:10px;font-weight:800;letter-spacing:1.5px}
.live-workspace-nav__title{color:#fff;font-size:20px;font-weight:800;letter-spacing:.2px}.live-workspace-nav__sub{margin-top:5px;color:#8baabd;font-size:11px}
.live-workspace-nav__back{display:flex;align-items:center;gap:8px;margin:16px 4px 20px;padding:10px 11px;border:1px solid rgba(135,202,221,.2);border-radius:8px;color:#b9d4df!important;font-size:12px;font-weight:650;text-decoration:none!important;transition:.18s ease}
.live-workspace-nav__back:hover{border-color:rgba(75,211,235,.48);background:rgba(68,201,229,.08);color:#fff!important}
.live-workspace-nav__label{margin:0 10px 8px;color:#6f94a7;font-size:10px;font-weight:750;letter-spacing:1px}
.live-workspace-nav__item{position:relative;display:block;margin:4px 0;padding:12px 13px 11px 16px;border-radius:8px;color:#b8d0dc!important;text-decoration:none!important;transition:.18s ease}
.live-workspace-nav__item:hover{background:rgba(83,206,232,.09);color:#fff!important}.live-workspace-nav__item.active{background:linear-gradient(90deg,rgba(31,193,221,.23),rgba(31,193,221,.08));color:#fff!important;box-shadow:inset 3px 0 0 #36d4eb}
.live-workspace-nav__name{display:block;font-size:13px;font-weight:760}.live-workspace-nav__help{display:block;margin-top:4px;color:#7196a9;font-size:10px;line-height:1.45}.live-workspace-nav__item.active .live-workspace-nav__help{color:#a8d5df}
section[data-testid="stSidebar"] div[data-testid="stButton"] button{min-height:40px!important;border:1px solid rgba(135,202,221,.20)!important;border-radius:8px!important;background:rgba(7,31,51,.55)!important;color:#c4dbe5!important;box-shadow:none!important}
section[data-testid="stSidebar"] div[data-testid="stButton"] button:hover{border-color:rgba(75,211,235,.50)!important;background:rgba(33,178,207,.13)!important;color:#fff!important}
section[data-testid="stSidebar"] div[data-testid="stButton"] button[kind="primary"]{border-color:rgba(54,212,235,.48)!important;background:linear-gradient(90deg,rgba(31,193,221,.28),rgba(31,193,221,.12))!important;color:#fff!important;box-shadow:inset 3px 0 0 #36d4eb!important}
section[data-testid="stSidebar"] div[data-testid="stButton"] button p{color:inherit!important;font-weight:700!important}
@media(max-width:900px){.main .block-container{padding:18px 18px 40px!important}}
.main div[data-testid="stVerticalBlock"]{gap:12px}
.main div[data-testid="stHorizontalBlock"]{gap:12px}
.main .page-hero{margin:0 0 16px!important;padding:20px 22px!important;border-radius:16px!important;box-shadow:0 6px 24px rgba(12,44,73,.055)!important}
.main .page-hero__title{font-size:26px!important}.main .page-hero__subtitle{font-size:13px!important}
.scope-note{margin:0 0 16px;padding:11px 14px;border:1px solid #d7e4ee;border-radius:10px;background:#f7fafc;color:#526579;font-size:12px;line-height:1.65}
.section-kicker{margin:0 0 2px;color:#10263d;font-size:14px;font-weight:760}.section-help{margin:0 0 10px;color:#6e7e8e;font-size:11px}
.main div[data-testid="stVerticalBlockBorderWrapper"]{border:1px solid #dfe7ee!important;border-radius:14px!important;background:#fff!important;box-shadow:0 6px 22px rgba(16,48,82,.045)!important}
.main div[data-testid="stVerticalBlockBorderWrapper"]>div{padding:16px!important}
.main div[data-testid="stMetric"]{min-height:98px!important;padding:14px 15px!important;border-radius:12px!important;box-shadow:0 4px 16px rgba(16,48,82,.045)!important}
.main div[data-testid="stMetricLabel"]{font-size:12px!important}.main div[data-testid="stMetricValue"]{font-size:23px!important}
.main div[data-testid="stTabs"] [data-baseweb="tab-list"]{width:100%!important;gap:3px!important;padding:4px!important;border-radius:11px!important;background:#eaf0f5!important}
.main div[data-testid="stTabs"] [data-baseweb="tab-list"]>div{overflow-x:auto!important;scrollbar-width:thin}
.main div[data-testid="stTabs"] button[role="tab"]{height:36px!important;padding:0 12px!important;border-radius:8px!important;font-size:12px!important;white-space:nowrap!important}
.main div[data-testid="stTabs"] button[role="tab"][aria-selected="true"]{color:#0a567d!important;box-shadow:0 2px 7px rgba(12,44,73,.1)!important}
.main [data-testid="stWidgetLabel"] p{font-size:11px!important;font-weight:650!important;color:#647589!important}
.main div[data-baseweb="select"]>div,.main div[data-baseweb="input"]>div{min-height:38px!important;border-radius:8px!important}
.main [data-testid="stSegmentedControl"]{margin-bottom:2px}.main [data-testid="stSegmentedControl"] button{min-height:34px!important;font-size:11px!important}
.main div[data-testid="stDataFrame"]{border-radius:11px!important;box-shadow:none!important;border-color:#dfe7ee!important}
.main div[data-testid="stPlotlyChart"]{padding:8px;border-radius:12px!important;box-shadow:none!important}
.main .stDownloadButton>button{height:38px!important;border-radius:8px!important;font-size:12px!important}
.main h3{margin-top:12px!important;margin-bottom:6px!important;font-size:17px!important}
.main [data-testid="stCaptionContainer"]{margin:2px 0 6px!important}
@media(max-width:900px){.main .block-container{padding:16px 12px 36px!important}.main .page-hero{padding:16px!important}}
</style>
<div class="scope-note">平台支付按直播发生时间统计；发货、退货按仓库发生时间统计。两种口径并列观察，不进行逐日强制对账。</div>
""", unsafe_allow_html=True)

today = date.today()
with st.container(border=True):
    st.markdown('<div class="section-kicker">分析周期</div><div class="section-help">选择快捷周期，或使用自定义日期。</div>', unsafe_allow_html=True)
    quick = st.segmented_control("分析周期", ["近7天", "近15天", "近30天", "本月", "自定义"], default="近30天", label_visibility="collapsed")
    if quick == "近7天":
        default_start = today - timedelta(days=6)
    elif quick == "近15天":
        default_start = today - timedelta(days=14)
    elif quick == "本月":
        default_start = today.replace(day=1)
    else:
        default_start = today - timedelta(days=29)

    date_cols = st.columns(2)
    with date_cols[0]:
        start_date = st.date_input("开始日期", default_start, disabled=quick != "自定义")
    with date_cols[1]:
        end_date = st.date_input("结束日期", today, disabled=quick != "自定义")
if start_date > end_date:
    st.error("开始日期不能晚于结束日期。")
    st.stop()

with st.spinner("正在加载直播场次……"):
    sessions = load_live_sessions(start_date, end_date)

if sessions.empty:
    st.warning("所选范围暂无直播记录。可调整日期，或先在系统设置上传直播数据。")
    st.stop()

sessions = sessions.copy()
sessions["start_time"] = pd.to_datetime(sessions["start_time"], errors="coerce")
sessions["直播日期"] = sessions["start_time"].dt.date

all_shops = sorted(sessions["shop_name"].dropna().astype(str).unique())
all_anchors = sorted(sessions["anchor_name"].dropna().astype(str).unique())
default_shop = "抖音引美25.7转愉燃轮廓轻奢旗舰店"
default_anchor = "轮廓官方旗舰店"
if "live_operations_shops" not in st.session_state:
    st.session_state["live_operations_shops"] = [default_shop] if default_shop in all_shops else all_shops
else:
    st.session_state["live_operations_shops"] = [
        value for value in st.session_state["live_operations_shops"] if value in all_shops
    ] or ([default_shop] if default_shop in all_shops else all_shops)
if "live_operations_anchors" not in st.session_state:
    st.session_state["live_operations_anchors"] = [default_anchor] if default_anchor in all_anchors else all_anchors
else:
    st.session_state["live_operations_anchors"] = [
        value for value in st.session_state["live_operations_anchors"] if value in all_anchors
    ] or ([default_anchor] if default_anchor in all_anchors else all_anchors)
with st.container(border=True):
    st.markdown('<div class="section-kicker">分析对象</div><div class="section-help">筛选需要比较的直播间与主播。</div>', unsafe_allow_html=True)
    filter_cols = st.columns(2)
    with filter_cols[0]:
        selected_shops = st.multiselect("直播间／店铺", all_shops, key="live_operations_shops")
    with filter_cols[1]:
        selected_anchors = st.multiselect("主播", all_anchors, key="live_operations_anchors")

sessions = sessions[sessions["shop_name"].isin(selected_shops) & sessions["anchor_name"].isin(selected_anchors)]
room_ids = sessions["live_room_id"].astype(str)
if sessions.empty:
    st.warning("当前筛选条件下没有直播记录。")
    st.stop()

room_ids_key = tuple(room_ids)
with st.spinner("正在加载所选直播间数据……"):
    products, metrics = load_live_room_data(room_ids_key)
    channels, talks = load_live_auxiliary(room_ids_key)
    actuals = load_live_actuals(start_date, end_date)
    product_master = load_product_master()

products = products.merge(
    sessions[["live_room_id", "shop_name", "anchor_name", "start_time", "直播日期", "duration_seconds"]],
    on="live_room_id", how="left", validate="many_to_one",
)
for column in ["paid_amount", "sold_units", "click_users", "talk_count", "pre_ship_refund_amount", "post_ship_refund_amount"]:
    products[column] = pd.to_numeric(products.get(column), errors="coerce").fillna(0)
if "product_image_url" not in products.columns:
    products["product_image_url"] = None

# 商品讲解区间：计算相对开播分钟、讲解效率和在线人数变化。
talk_summary = pd.DataFrame(columns=[
    "style_code", "讲解场次", "讲解区间数", "累计讲解分钟", "讲解区间支付",
    "讲解区间成交件数", "在线净变化", "平均在线变化", "在线上升场次占比", "分均在线人数",
])
if not talks.empty:
    talks["style_code"] = talks["style_code"].fillna("").astype(str).str.strip().str.upper()
    talks = talks[talks["style_code"] != ""].copy()
    talk_session_times = sessions[["live_room_id", "start_time"]].copy()
    talk_session_times["live_room_id"] = talk_session_times["live_room_id"].astype(str)
    talk_session_times["session_start_utc"] = pd.to_datetime(talk_session_times["start_time"], utc=True, errors="coerce")
    talks["live_room_id"] = talks["live_room_id"].astype(str)
    talks = talks.merge(talk_session_times[["live_room_id", "session_start_utc"]], on="live_room_id", how="left")
    talks["talk_start_time"] = parse_epoch_utc(talks["talk_start_epoch"])
    talks["talk_end_time"] = parse_epoch_utc(talks["talk_end_epoch"])
    talks["开播后分钟"] = ((talks["talk_start_time"] - talks["session_start_utc"]).dt.total_seconds() / 60).clip(lower=0)
    for column in ["talk_duration_seconds", "paid_amount", "sold_units", "viewer_change", "avg_online_users"]:
        talks[column] = pd.to_numeric(talks.get(column), errors="coerce").fillna(0)
    talks["讲解分钟"] = talks["talk_duration_seconds"] / 60
    talks["在线上升"] = talks["viewer_change"] > 0
    talk_summary = talks.groupby("style_code", as_index=False).agg(
        讲解场次=("live_room_id", "nunique"), 讲解区间数=("product_id", "size"),
        累计讲解分钟=("讲解分钟", "sum"), 讲解区间支付=("paid_amount", "sum"),
        讲解区间成交件数=("sold_units", "sum"), 在线净变化=("viewer_change", "sum"),
        平均在线变化=("viewer_change", "mean"), 在线上升场次占比=("在线上升", "mean"),
        分均在线人数=("avg_online_users", "mean"),
    )
    talk_summary["支付/讲解分钟"] = talk_summary["讲解区间支付"].div(talk_summary["累计讲解分钟"].replace(0, pd.NA))
    talk_summary["成交件数/讲解分钟"] = talk_summary["讲解区间成交件数"].div(talk_summary["累计讲解分钟"].replace(0, pd.NA))
    talk_summary["每次讲解平均产出"] = talk_summary["讲解区间支付"].div(talk_summary["讲解区间数"].replace(0, pd.NA))

actual_scope = "店铺＋主播＋货号"
actual_by_style = pd.DataFrame(columns=["style_code", "ship_amount", "return_amount", "net_amount"])
if not actuals.empty and "style_code" in actuals:
    actuals = actuals.copy()
    actuals["shop_key"] = actuals["shop_name"].fillna("").astype(str).str.strip().str.upper()
    actuals = actuals[actuals["shop_key"].isin({x.strip().upper() for x in selected_shops})]
    anchor_source = actuals.get("anchor", actuals.get("anchor_display", pd.Series(index=actuals.index, dtype=str)))
    actuals["anchor_key"] = anchor_source.fillna("").astype(str).str.strip().str.upper()
    anchor_match = actuals["anchor_key"].isin({x.strip().upper() for x in selected_anchors})
    if anchor_match.any():
        actuals = actuals[anchor_match]
    else:
        actual_scope = "店铺＋货号（无法可靠区分主播）"
    actuals["style_code"] = actuals["style_code"].astype(str).str.strip().str.upper()
    aggregations = {
        "ship_amount": ("ship_amount", "sum"),
        "return_amount": ("return_amount", "sum"),
        "net_amount": ("net_amount", "sum"),
    }
    if "brand" in actuals.columns:
        aggregations["销售品牌"] = (
            "brand",
            lambda values: next(
                (str(value).strip() for value in values if pd.notna(value) and str(value).strip()),
                None,
            ),
        )
    actual_by_style = actuals.groupby("style_code", as_index=False).agg(**aggregations)

style_summary = products[products["style_code"].notna()].groupby("style_code", as_index=False).agg(
    商品名称=("product_name", "first"), 商品图片=("product_image_url", "first"),
    上播场次=("live_room_id", "nunique"),
    讲解次数=("talk_count", "sum"), 累计点击=("click_users", "sum"),
    成交件数=("sold_units", "sum"), 平台支付=("paid_amount", "sum"),
    平台退款=("pre_ship_refund_amount", "sum"),
)
style_summary = style_summary.merge(actual_by_style, on="style_code", how="left").fillna(0)
style_summary = style_summary.merge(talk_summary, on="style_code", how="left")
if not product_master.empty and "style_code" in product_master.columns:
    master_meta = product_master.copy()
    master_meta["style_code"] = master_meta["style_code"].astype(str).str.strip().str.upper()
    if "category" in master_meta.columns:
        master_meta["品类"] = master_meta["category"]
    elif "master_category" in master_meta.columns:
        master_meta["品类"] = master_meta["master_category"]
    elif "product_category" in master_meta.columns:
        master_meta["品类"] = master_meta["product_category"]
    else:
        master_meta["品类"] = None
    brand_source = master_meta.get("brand", pd.Series(index=master_meta.index, dtype="object"))
    year_source = master_meta.get(
        "product_year",
        master_meta.get("year", pd.Series(index=master_meta.index, dtype="object")),
    )
    master_meta["品牌"] = brand_source
    master_meta["年份"] = year_source
    master_meta["资料库图片"] = master_meta["image_url"] if "image_url" in master_meta.columns else None
    master_meta["上新日期"] = master_meta["launch_date"] if "launch_date" in master_meta.columns else None
    master_meta = master_meta[["style_code", "品牌", "年份", "品类", "上新日期", "资料库图片"]].drop_duplicates("style_code")
    style_summary = style_summary.merge(master_meta, on="style_code", how="left")
    parsed_brand = style_summary["style_code"].astype(str).str.slice(0, 1)
    parsed_year = style_summary["style_code"].astype(str).str.slice(1, 3)
    master_brand = style_summary["品牌"].replace([0, "0", "", "nan", "None"], pd.NA)
    sales_brand = style_summary.get("销售品牌", pd.Series(index=style_summary.index, dtype="object")).replace(
        [0, "0", "", "nan", "None"], pd.NA
    )
    style_summary["品牌"] = master_brand.fillna(sales_brand).fillna(parsed_brand)
    style_summary["年份"] = style_summary["年份"].replace(
        [0, "0", "", "nan", "None"], pd.NA
    ).fillna(parsed_year)
    style_summary["商品图片"] = style_summary["商品图片"].replace([0, "0", ""], pd.NA).fillna(style_summary["资料库图片"])
    style_summary = style_summary.drop(columns="资料库图片")
else:
    parsed_brand = style_summary["style_code"].astype(str).str.slice(0, 1)
    parsed_year = style_summary["style_code"].astype(str).str.slice(1, 3)
    sales_brand = style_summary.get("销售品牌", pd.Series(index=style_summary.index, dtype="object")).replace(
        [0, "0", "", "nan", "None"], pd.NA
    )
    style_summary["品牌"] = sales_brand.fillna(parsed_brand)
    style_summary["年份"] = parsed_year
    style_summary["品类"] = None
style_summary = style_summary.drop(columns="销售品牌", errors="ignore")
style_image_lookup = (
    style_summary.assign(style_code=style_summary["style_code"].astype(str).str.strip().str.upper())
    .set_index("style_code")["商品图片"].dropna().to_dict()
)
talk_columns = [
    "讲解场次", "讲解区间数", "累计讲解分钟", "讲解区间支付", "讲解区间成交件数",
    "在线净变化", "平均在线变化", "在线上升场次占比", "分均在线人数",
    "支付/讲解分钟", "成交件数/讲解分钟", "每次讲解平均产出",
]
for column in talk_columns:
    if column in style_summary:
        style_summary[column] = pd.to_numeric(style_summary[column], errors="coerce").fillna(0)
    else:
        style_summary[column] = 0.0
style_summary["点击成交率"] = style_summary["成交件数"].div(style_summary["累计点击"].replace(0, pd.NA))
style_summary["成交场次率"] = products.assign(成交=products["sold_units"] > 0).groupby("style_code")["成交"].mean().reindex(style_summary["style_code"]).values
style_summary["退款率"] = style_summary["平台退款"].div(style_summary["平台支付"].replace(0, pd.NA))

# 稳定复销不是单纯的成交场次率排行：同时判断跨周持续性、单场依赖、近期表现和退款健康度。
session_style = products[products["style_code"].notna()].groupby(
    ["style_code", "live_room_id", "直播日期"], as_index=False
).agg(场次成交件数=("sold_units", "sum"), 场次支付=("paid_amount", "sum"))
session_style["有成交"] = session_style["场次成交件数"] > 0
session_style["自然周"] = pd.to_datetime(session_style["直播日期"]).dt.to_period("W-SAT").astype(str)
stability_records = []
for style_code, history in session_style.groupby("style_code"):
    history = history.sort_values("直播日期")
    session_count = history["live_room_id"].nunique()
    sold_sessions = int(history["有成交"].sum())
    sold_units = history["场次成交件数"].sum()
    top_session_share = history["场次成交件数"].max() / sold_units if sold_units else 0
    recent_three_sales = int(history.tail(3)["有成交"].sum())
    stability_records.append({
        "style_code": style_code,
        "复销上播场次": session_count,
        "复销成交场次": sold_sessions,
        "复销成交场次率": sold_sessions / session_count if session_count else 0,
        "成交自然周数": history.loc[history["有成交"], "自然周"].nunique(),
        "最高单场成交占比": top_session_share,
        "最近3场成交场次": recent_three_sales,
    })
style_summary = style_summary.merge(pd.DataFrame(stability_records), on="style_code", how="left")
if "品类" in style_summary.columns:
    category_refund = style_summary.groupby("品类", dropna=False)["退款率"].transform("mean")
else:
    category_refund = pd.Series(style_summary["退款率"].mean(), index=style_summary.index)
style_summary["同品类平均退款率"] = category_refund.fillna(style_summary["退款率"].mean()).fillna(0)

def classify_repeat_sales(row):
    sessions_count = int(row.get("复销上播场次", 0) or 0)
    rate = float(row.get("复销成交场次率", 0) or 0)
    weeks = int(row.get("成交自然周数", 0) or 0)
    top_share = float(row.get("最高单场成交占比", 0) or 0)
    recent = int(row.get("最近3场成交场次", 0) or 0)
    clicks = float(row.get("累计点击", 0) or 0)
    sold_sessions = int(row.get("复销成交场次", 0) or 0)
    sold_units = float(row.get("成交件数", 0) or 0)
    refund = row.get("退款率")
    refund_benchmark = float(row.get("同品类平均退款率", 0) or 0)
    refund_risk = not pd.isna(refund) and float(refund) > refund_benchmark
    risk_note = "；退款率高于同品类平均" if refund_risk else ""
    if sessions_count < 2 or clicks < 20:
        return "样本不足", "上播少于2场或累计点击少于20"
    if sold_sessions >= 3 and recent == 0:
        return "近期转弱", f"历史至少3场成交，但最近3场均未成交{risk_note}"
    if sold_units >= 5 and top_share >= 0.7:
        return "单场爆发", f"累计成交至少5件，且最高单场贡献达到{top_share:.2%}{risk_note}"
    if sessions_count >= 5 and sold_sessions >= 3 and rate >= 0.6 and weeks >= 2 and recent >= 2:
        return "高置信稳定", f"至少5场、3场成交、成交率≥60%、跨2周且最近3场≥2场成交{risk_note}"
    if sessions_count >= 3 and sold_sessions >= 2 and rate >= 0.5 and recent >= 1:
        return "稳定复销", f"至少3场、2场成交、成交率≥50%且最近3场仍有成交{risk_note}"
    if sessions_count >= 2 and sold_sessions >= 1:
        return "有复销潜力", f"至少上播2场并产生过成交，继续积累重复成交样本{risk_note}"
    return "待观察", f"已有一定样本，但尚未形成重复成交{risk_note}"

repeat_classification = style_summary.apply(classify_repeat_sales, axis=1, result_type="expand")
style_summary[["复销分级", "复销判断依据"]] = repeat_classification

if not metrics.empty:
    duration_mask = metrics["metric_name"].astype(str).eq("人均观看时长")
    duration_numeric = pd.to_numeric(metrics.loc[duration_mask, "metric_value"], errors="coerce")
    needs_duration_recovery = duration_numeric.isna() | duration_numeric.eq(0)
    recovery_index = duration_numeric.index[needs_duration_recovery]
    metrics.loc[recovery_index, "metric_value"] = metrics.loc[recovery_index, "raw_value"].map(
        _duration_seconds
    )

metric_wide = metrics.pivot_table(
    index="live_room_id", columns="metric_name", values="metric_value", aggfunc="max"
) if not metrics.empty else pd.DataFrame()

# 场次经营链路：同名指标跨分组取最大值，避免“新增粉丝数”等重复累计。
session_kpis = sessions[[
    "live_room_id", "shop_name", "anchor_name", "start_time", "直播日期", "duration_seconds"
]].copy()
if not metric_wide.empty:
    metric_for_join = metric_wide.reset_index()
    metric_for_join["live_room_id"] = metric_for_join["live_room_id"].astype(str)
    session_kpis["live_room_id"] = session_kpis["live_room_id"].astype(str)
    session_kpis = session_kpis.merge(metric_for_join, on="live_room_id", how="left")

session_product = products.groupby("live_room_id", as_index=False).agg(
    商品明细点击人数=("click_users", "sum"), 商品成交件数=("sold_units", "sum"),
    商品支付金额=("paid_amount", "sum"),
)
session_product["live_room_id"] = session_product["live_room_id"].astype(str)
session_kpis = session_kpis.merge(session_product, on="live_room_id", how="left")

required_session_metrics = [
    "直播间曝光人数", "直播间观看人数", "平均在线人数", "最高在线人数",
    "自然流量观看人数", "付费流量观看人数", "直播间成交人数", "新增粉丝数",
    "人均观看时长", "点赞次数", "评论次数", "直播间用户支付金额",
    "直播间观看-互动率(人数)",
    "千次观看用户支付金额", "投放消耗（店铺绑定）", "投放消耗（店铺被投）",
    "首购率", "粉丝成交人数占比", "粉丝用户支付金额占比",
    "商品点击人数", "商品成交件数", "商品支付金额",
]
for column in required_session_metrics:
    if column not in session_kpis:
        session_kpis[column] = 0.0
    session_kpis[column] = pd.to_numeric(session_kpis[column], errors="coerce").fillna(0)
session_kpis["商品点击人数"] = session_kpis["商品明细点击人数"].where(
    session_kpis["商品明细点击人数"] > 0, session_kpis["商品点击人数"]
)

session_kpis["直播小时"] = session_kpis["duration_seconds"].div(3600).replace(0, pd.NA)
session_kpis["曝光进入率"] = session_kpis["直播间观看人数"].div(session_kpis["直播间曝光人数"].replace(0, pd.NA))
session_kpis["商品点击率"] = session_kpis["商品点击人数"].div(session_kpis["直播间观看人数"].replace(0, pd.NA))
session_kpis["观看成交率"] = session_kpis["直播间成交人数"].div(session_kpis["直播间观看人数"].replace(0, pd.NA))
session_kpis["点击成交率"] = session_kpis["直播间成交人数"].div(session_kpis["商品点击人数"].replace(0, pd.NA))
session_kpis["关注转化率"] = session_kpis["新增粉丝数"].div(session_kpis["直播间观看人数"].replace(0, pd.NA))
session_kpis["计算互动率"] = (session_kpis["点赞次数"] + session_kpis["评论次数"]).div(
    session_kpis["直播间观看人数"].replace(0, pd.NA)
)
session_kpis["互动率"] = session_kpis["直播间观看-互动率(人数)"].where(
    session_kpis["直播间观看-互动率(人数)"] > 0, session_kpis["计算互动率"]
)
session_kpis["自然流量占比"] = session_kpis["自然流量观看人数"].div(
    (session_kpis["自然流量观看人数"] + session_kpis["付费流量观看人数"]).replace(0, pd.NA)
)
session_kpis["每小时观看人数"] = session_kpis["直播间观看人数"].div(session_kpis["直播小时"])
session_kpis["每小时成交人数"] = session_kpis["直播间成交人数"].div(session_kpis["直播小时"])
session_kpis["每小时新增粉丝"] = session_kpis["新增粉丝数"].div(session_kpis["直播小时"])
session_kpis["每小时支付"] = session_kpis["商品支付金额"].div(session_kpis["直播小时"])
session_kpis["投放消耗观察值"] = session_kpis[[
    "投放消耗（店铺绑定）", "投放消耗（店铺被投）"
]].max(axis=1)
session_kpis["ROI"] = session_kpis["商品支付金额"].div(
    session_kpis["投放消耗观察值"].replace(0, pd.NA)
)

live_sections = {
    "决策总览": "先看经营问题、机会和下一场建议",
    "开播选品": "安排开场、极速流和稳定复销商品",
    "单场复盘": "复盘单场链路并形成下场动作",
    "商品决策": "查看单款表现、实销和主播适配",
    "主播对比": "比较主播效率及同商品表现",
}
if st.session_state.get("live_operations_section") not in live_sections:
    st.session_state["live_operations_section"] = "决策总览"

active_section = st.session_state["live_operations_section"]


def set_live_operations_section(section_name: str) -> None:
    st.session_state["live_operations_section"] = section_name


with st.sidebar:
    st.markdown(
        f'''<nav class="live-workspace-nav">
            <div class="live-workspace-nav__brand">
                <div class="live-workspace-nav__eyebrow">LIVE OPERATIONS</div>
                <div class="live-workspace-nav__title">直播经营工作台</div>
                <div class="live-workspace-nav__sub">从数据表现走向下一场动作</div>
            </div>
        </nav>''',
        unsafe_allow_html=True,
    )
    if st.button("← 返回数据罗盘", key="back_to_data_compass", width="stretch"):
        st.switch_page("pages/dashboard.py")
    st.markdown('<div class="live-workspace-nav__label">分析任务</div>', unsafe_allow_html=True)
    for section_name, section_help in live_sections.items():
        st.button(
            section_name,
            key=f"live_section_{section_name}",
            help=section_help,
            type="primary" if section_name == active_section else "secondary",
            width="stretch",
            on_click=set_live_operations_section,
            args=(section_name,),
        )
st.markdown(
    f'<div class="section-kicker">{active_section}</div><div class="section-help">{live_sections[active_section]}</div>',
    unsafe_allow_html=True,
)

hidden_section_holders = []


def section_target(section_name):
    if active_section == section_name:
        return st.container()
    holder = st.empty()
    hidden_section_holders.append(holder)
    return holder.container()


decision_tab = section_target("决策总览")
selection_tab = section_target("开播选品")
review_tab = section_target("单场复盘")
product_tab = section_target("商品决策")
compare_tab = section_target("主播对比")

with decision_tab:
    st.subheader("本周期经营结论")
    decision_records = []
    total_exposure = session_kpis["直播间曝光人数"].sum()
    total_viewers = session_kpis["直播间观看人数"].sum()
    total_clicks = session_kpis["商品点击人数"].sum()
    total_buyers = session_kpis["直播间成交人数"].sum()
    total_follows = session_kpis["新增粉丝数"].sum()
    total_spend = session_kpis["投放消耗观察值"].sum()
    total_platform_pay = session_kpis["商品支付金额"].sum()
    weighted_watch_seconds = (
        (session_kpis["人均观看时长"] * session_kpis["直播间观看人数"]).sum() / total_viewers
        if total_viewers else 0
    )
    overall_entry_rate = total_viewers / total_exposure if total_exposure else 0
    overall_click_rate = total_clicks / total_viewers if total_viewers else 0
    overall_click_conversion = total_buyers / total_clicks if total_clicks else 0
    overall_follow_rate = total_follows / total_viewers if total_viewers else 0
    overall_spend_output = total_platform_pay / total_spend if total_spend else None

    st.markdown("#### 直播间经营链路")
    funnel_cards = st.columns(6)
    funnel_cards[0].metric("曝光进入率", f"{overall_entry_rate:.2%}")
    funnel_cards[1].metric("人均观看时长", f"{weighted_watch_seconds:,.0f}秒")
    funnel_cards[2].metric("商品点击率", f"{overall_click_rate:.2%}")
    funnel_cards[3].metric("点击成交率", f"{overall_click_conversion:.2%}")
    funnel_cards[4].metric("关注转化率", f"{overall_follow_rate:.2%}")
    funnel_cards[5].metric("ROI", f"{overall_spend_output:.2f}" if overall_spend_output is not None else "—")
    st.caption("ROI为经营观察指标：平台支付金额 ÷ 投放消耗观察值；投放消耗取“店铺绑定”和“店铺被投”两者较大值。")

    valid_rooms = session_kpis[session_kpis["直播间观看人数"] > 0].copy()
    non_product_rules = [
        ("曝光承接偏弱", "曝光进入率", "直播间曝光人数", "优化封面、标题、开场内容和进房承接", False),
        ("停留偏弱", "人均观看时长", "直播间观看人数", "优化开场节奏和内容钩子，减少无效停顿", False),
        ("商品点击偏弱", "商品点击率", "商品点击人数", "加强商品利益点、挂车引导和展示节奏", False),
        ("成交承接偏弱", "点击成交率", "直播间成交人数", "检查价格机制、信任表达和成交话术", False),
        ("粉丝沉淀偏弱", "关注转化率", "新增粉丝数", "增加明确关注理由和粉丝权益表达", False),
        ("互动偏弱", "互动率", "直播间观看人数", "增加互动问题、场景引导和节奏节点", False),
        ("新客成交偏弱", "首购率", "直播间成交人数", "检查新客利益机制和首次购买信任表达", False),
        ("付费流量效率偏弱", "ROI", "投放消耗观察值", "复盘投放流量进入后的停留、点击和成交承接", True),
    ]
    for rule_name, metric_column, sample_column, action, needs_spend in non_product_rules:
        comparable = valid_rooms[valid_rooms[sample_column] > 0].copy()
        if needs_spend:
            comparable = comparable[comparable["投放消耗观察值"] > 0]
        median_value = comparable[metric_column].median() if not comparable.empty else pd.NA
        if pd.isna(median_value) or median_value <= 0:
            continue
        weak_rooms = comparable[comparable[metric_column] < median_value * .7]
        if weak_rooms.empty:
            continue
        weakest = weak_rooms.sort_values(metric_column).iloc[0]
        metric_value = safe_number(weakest.get(metric_column))
        metric_display = f"{metric_value:.2%}" if "率" in metric_column else f"{metric_value:.2f}"
        decision_records.append({
            "优先级": "高" if len(weak_rooms) >= max(2, len(comparable) // 3) else "中",
            "类型": rule_name, "货号": "—", "商品名称": "直播间整体",
            "发现": f"{len(weak_rooms)}场低于场次中位水平30%以上；最低{metric_column}{metric_display}",
            "建议动作": action, "可信度": "高可信" if len(comparable) >= 5 else "中可信",
            "上播场次": len(comparable),
        })

    median_talk_output = style_summary.loc[
        style_summary["累计讲解分钟"] > 0, "支付/讲解分钟"
    ].median() if "支付/讲解分钟" in style_summary else 0
    median_talk_output = float(median_talk_output) if pd.notna(median_talk_output) else 0
    for _, row in style_summary.iterrows():
        code = safe_text(row.get("style_code"))
        name = safe_text(row.get("商品名称"))
        sessions_count = int(safe_number(row.get("上播场次")))
        clicks = safe_number(row.get("累计点击"))
        conversion = safe_number(row.get("点击成交率"))
        refund_rate = safe_number(row.get("退款率"))
        repeat_level = safe_text(row.get("复销分级"))
        talk_minutes = safe_number(row.get("累计讲解分钟"))
        talk_output = safe_number(row.get("支付/讲解分钟"))
        online_change = safe_number(row.get("平均在线变化"))
        talk_intervals = int(safe_number(row.get("讲解区间数")))
        confidence = "高可信" if sessions_count >= 5 else "中可信" if sessions_count >= 3 else "待验证"
        if clicks >= 100 and conversion < .03:
            decision_records.append({
                "优先级": "高", "类型": "高点击低成交", "货号": code, "商品名称": name,
                "发现": f"累计点击{clicks:,.0f}，点击成交率仅{conversion:.2%}",
                "建议动作": "下一场缩短讲解并重点测试价格、利益点和尺码说明",
                "可信度": confidence, "上播场次": sessions_count,
            })
        if refund_rate >= .35:
            decision_records.append({
                "优先级": "高", "类型": "退款风险", "货号": code, "商品名称": name,
                "发现": f"平台退款率达到{refund_rate:.2%}",
                "建议动作": "核查退款原因，确认商品问题前暂缓作为主推款",
                "可信度": confidence, "上播场次": sessions_count,
            })
        if repeat_level in {"高置信稳定", "稳定复销"}:
            decision_records.append({
                "优先级": "机会", "类型": "稳定复销", "货号": code, "商品名称": name,
                "发现": f"{repeat_level}，成交场次率{safe_number(row.get('复销成交场次率')):.2%}",
                "建议动作": "加入下一场常规排品，并优先放入前60分钟测试",
                "可信度": "高可信" if repeat_level == "高置信稳定" else confidence,
                "上播场次": sessions_count,
            })
        if talk_intervals >= 3 and talk_minutes >= 15 and median_talk_output and talk_output < median_talk_output * .5:
            decision_records.append({
                "优先级": "中", "类型": "讲解效率偏低", "货号": code, "商品名称": name,
                "发现": f"累计讲解{talk_minutes:.1f}分钟，分钟产出¥{talk_output:,.0f}",
                "建议动作": "减少单次讲解时长，调整讲解顺序后再测试",
                "可信度": confidence, "上播场次": sessions_count,
            })
        if talk_intervals >= 3 and online_change < 0:
            decision_records.append({
                "优先级": "中", "类型": "在线人数流失", "货号": code, "商品名称": name,
                "发现": f"讲解期间平均在线人数变化{online_change:+.1f}",
                "建议动作": "避免放在流量高峰，先调整话术或搭配方式",
                "可信度": confidence, "上播场次": sessions_count,
            })

    decisions = pd.DataFrame(decision_records)
    priority_order = {"高": 0, "中": 1, "机会": 2}
    if not decisions.empty:
        decisions["优先级排序"] = decisions["优先级"].map(priority_order).fillna(9)
        decisions = decisions.sort_values(["优先级排序", "上播场次"], ascending=[True, False]).drop(columns="优先级排序")
        issue_count = int(decisions["优先级"].isin(["高", "中"]).sum())
        opportunity_count = int((decisions["优先级"] == "机会").sum())
    else:
        issue_count = opportunity_count = 0

    decision_cards = st.columns(4)
    decision_cards[0].metric("待处理问题", f"{issue_count}项")
    decision_cards[1].metric("可放大机会", f"{opportunity_count}项")
    decision_cards[2].metric("稳定复销商品", f"{(style_summary['复销分级'].isin(['高置信稳定', '稳定复销'])).sum()}款")
    decision_cards[3].metric("高点击低成交", f"{((style_summary['累计点击'] >= 100) & (style_summary['点击成交率'] < .03)).sum()}款")

    if decisions.empty:
        st.info("当前周期暂未识别到达到规则阈值的问题或机会，可扩大日期范围继续观察。")
    else:
        st.markdown("#### 优先处理清单")
        show_table(decisions.head(12))
        evidence_options = [
            code for code in decisions["货号"].drop_duplicates().tolist() if str(code).strip() not in {"", "—"}
        ]
        if evidence_options:
            evidence_code = st.selectbox(
                "查看商品判断证据", evidence_options,
                format_func=lambda code: f"{code}｜{decisions.loc[decisions['货号'] == code, '商品名称'].iloc[0]}",
                key="decision_evidence_style",
            )
            evidence_rows = session_style[session_style["style_code"].astype(str) == str(evidence_code)].merge(
                sessions[["live_room_id", "shop_name", "anchor_name", "start_time"]],
                on="live_room_id", how="left",
            )
            with st.expander("展开商品逐场证据"):
                show_table(evidence_rows.sort_values("start_time", ascending=False))
        with st.expander("展开直播间逐场经营链路"):
            show_table(session_kpis[[
                "start_time", "shop_name", "anchor_name", "直播间曝光人数", "直播间观看人数",
                "曝光进入率", "人均观看时长", "平均在线人数", "最高在线人数", "自然流量观看人数",
                "付费流量观看人数", "互动率", "商品点击率", "点击成交率", "新增粉丝数", "关注转化率",
                "首购率", "粉丝成交人数占比", "粉丝用户支付金额占比",
                "千次观看用户支付金额", "投放消耗观察值", "ROI",
            ]].sort_values("start_time", ascending=False))

    st.markdown("#### 下一场建议优先测试")
    next_session_candidates = style_summary[
        style_summary["复销分级"].isin(["高置信稳定", "稳定复销", "有复销潜力"])
    ].copy()
    if not next_session_candidates.empty:
        next_session_candidates["建议分"] = (
            next_session_candidates["复销成交场次率"].fillna(0) * 45
            + next_session_candidates["点击成交率"].fillna(0).clip(upper=.2) / .2 * 25
            + next_session_candidates["在线上升场次占比"].fillna(0) * 15
            + (1 - next_session_candidates["退款率"].fillna(0).clip(upper=1)) * 15
        )
        next_display = next_session_candidates.sort_values("建议分", ascending=False).head(8)[[
            "style_code", "商品名称", "复销分级", "上播场次", "成交场次率", "点击成交率",
            "支付/讲解分钟", "在线上升场次占比", "退款率", "建议分",
        ]]
        show_table(next_display)
        st.caption("建议分用于排序，不作为业务目标；由复销稳定性、转化、在线变化和退款健康度共同组成。")
    else:
        st.info("当前还没有达到复销候选条件的商品。")

    st.markdown("#### 经营结果概览")
    total_ship = actual_by_style["ship_amount"].sum()
    total_return = actual_by_style["return_amount"].sum()
    cards = st.columns(5)
    cards[0].metric("历史有效场次", f"{sessions['live_room_id'].nunique():,}")
    cards[1].metric("直播总时长", f"{sessions['duration_seconds'].sum()/3600:,.1f}小时")
    cards[2].metric("平台支付", f"¥{products['paid_amount'].sum():,.0f}")
    cards[3].metric("范围实销", f"¥{(total_ship-total_return):,.0f}")
    cards[4].metric("历史商品样本", f"{len(products):,}")
    st.caption(f"范围实销关联层级：{actual_scope}；实销＝发货－退货，不归因为单场直播。")
    daily_platform = products.groupby("直播日期", as_index=False)["paid_amount"].sum().rename(columns={"paid_amount": "平台支付"})
    if not actuals.empty:
        actual_daily = actuals.groupby("sale_date", as_index=False).agg(发货=("ship_amount", "sum"), 退货=("return_amount", "sum"), 实销=("net_amount", "sum"))
        actual_daily["直播日期"] = pd.to_datetime(actual_daily["sale_date"]).dt.date
        daily_platform = daily_platform.merge(actual_daily.drop(columns="sale_date"), on="直播日期", how="outer").fillna(0)
    trend_columns = [x for x in ["平台支付", "发货", "退货", "实销"] if x in daily_platform]
    st.plotly_chart(px.line(daily_platform.sort_values("直播日期"), x="直播日期", y=trend_columns, markers=True, title="历史经营趋势"), width="stretch")

with review_tab:
    labels = sessions.sort_values("start_time", ascending=False).assign(
        场次=lambda frame: frame["start_time"].dt.strftime("%m-%d %H:%M") + "｜" + frame["shop_name"] + "｜" + frame["anchor_name"]
    )
    selected_label = st.selectbox("选择直播场次", labels["场次"].tolist())
    room = str(labels.loc[labels["场次"] == selected_label, "live_room_id"].iloc[0])
    room_products = products[products["live_room_id"].astype(str) == room].copy()
    room_metrics = metric_wide.loc[room] if room in metric_wide.index.astype(str) else pd.Series(dtype=float)
    cards = st.columns(5)
    cards[0].metric("观看人数", f"{float(room_metrics.get('直播间观看人数', 0) or 0):,.0f}")
    cards[1].metric("商品点击人数", f"{room_products['click_users'].sum():,.0f}")
    cards[2].metric("成交人数", f"{float(room_metrics.get('直播间成交人数', 0) or 0):,.0f}")
    cards[3].metric("成交件数", f"{room_products['sold_units'].sum():,.0f}")
    cards[4].metric("平台支付", f"¥{room_products['paid_amount'].sum():,.0f}")
    room_kpi_rows = session_kpis[session_kpis["live_room_id"].astype(str) == room]
    room_kpi = room_kpi_rows.iloc[0] if not room_kpi_rows.empty else pd.Series(dtype=float)
    st.markdown("#### 本场经营链路")
    flow_cards = st.columns(6)
    flow_cards[0].metric("曝光进入率", f"{safe_number(room_kpi.get('曝光进入率')):.2%}")
    flow_cards[1].metric("人均观看时长", f"{safe_number(room_kpi.get('人均观看时长')):,.0f}秒")
    flow_cards[2].metric("商品点击率", f"{safe_number(room_kpi.get('商品点击率')):.2%}")
    flow_cards[3].metric("点击成交率", f"{safe_number(room_kpi.get('点击成交率')):.2%}")
    flow_cards[4].metric("关注转化率", f"{safe_number(room_kpi.get('关注转化率')):.2%}")
    room_spend_output = room_kpi.get("ROI", pd.NA)
    flow_cards[5].metric("ROI", f"{safe_number(room_spend_output):.2f}" if pd.notna(room_spend_output) else "—")
    audience_cards = st.columns(6)
    audience_cards[0].metric("自然／付费观看", f"{safe_number(room_kpi.get('自然流量观看人数')):,.0f}／{safe_number(room_kpi.get('付费流量观看人数')):,.0f}")
    audience_cards[1].metric("新增粉丝", f"{safe_number(room_kpi.get('新增粉丝数')):,.0f}")
    audience_cards[2].metric("首购率", f"{safe_number(room_kpi.get('首购率')):.2%}")
    audience_cards[3].metric("粉丝成交人数占比", f"{safe_number(room_kpi.get('粉丝成交人数占比')):.2%}")
    audience_cards[4].metric("千次观看支付", f"¥{safe_number(room_kpi.get('千次观看用户支付金额')):,.0f}")
    audience_cards[5].metric("互动率", f"{safe_number(room_kpi.get('互动率')):.2%}")
    st.markdown("#### 本场诊断与下场动作")
    room_actions = []
    room_benchmarks = {
        column: valid_rooms.loc[valid_rooms[column].notna(), column].median()
        for column in ["曝光进入率", "人均观看时长", "互动率", "商品点击率", "点击成交率", "关注转化率", "首购率", "ROI"]
    }
    room_non_product_rules = [
        ("曝光进入率", "曝光承接", "优化封面、标题和开场进房承接"),
        ("人均观看时长", "停留承接", "优化开场节奏和内容钩子，减少无效停顿"),
        ("互动率", "互动承接", "增加互动问题、场景引导和节奏节点"),
        ("商品点击率", "商品点击", "加强利益点表达、挂车引导和商品展示"),
        ("点击成交率", "成交承接", "检查价格机制、信任表达和逼单节奏"),
        ("关注转化率", "粉丝沉淀", "增加明确关注理由和粉丝权益表达"),
        ("首购率", "新客成交", "检查新客利益机制和首次购买信任表达"),
        ("ROI", "付费流量效率", "复盘投放流量进入后的停留和成交承接"),
    ]
    for metric_column, diagnosis_name, action in room_non_product_rules:
        value = room_kpi.get(metric_column, pd.NA)
        benchmark = room_benchmarks.get(metric_column, pd.NA)
        if pd.isna(value) or pd.isna(benchmark) or benchmark <= 0 or value >= benchmark * .7:
            continue
        value_display = f"{safe_number(value):.2%}" if "率" in metric_column else f"{safe_number(value):.2f}"
        benchmark_display = f"{safe_number(benchmark):.2%}" if "率" in metric_column else f"{safe_number(benchmark):.2f}"
        room_actions.append({
            "优先级": "高", "本场发现": f"{diagnosis_name}偏弱：{metric_column}{value_display}，场次中位值{benchmark_display}",
            "下场动作": action,
        })
    room_clicks = float(room_products["click_users"].sum())
    room_units = float(room_products["sold_units"].sum())
    room_conversion = room_units / room_clicks if room_clicks else 0
    room_refunds = float(room_products["pre_ship_refund_amount"].sum() + room_products["post_ship_refund_amount"].sum())
    room_paid = float(room_products["paid_amount"].sum())
    room_refund_rate = room_refunds / room_paid if room_paid else 0
    if room_clicks >= 100 and room_conversion < .03:
        room_actions.append({
            "优先级": "高", "本场发现": f"点击成交率仅{room_conversion:.2%}",
            "下场动作": "优先复盘价格机制、利益点表达和尺码说明，不直接增加讲解时长",
        })
    if room_refund_rate >= .35:
        room_actions.append({
            "优先级": "高", "本场发现": f"平台退款占支付{room_refund_rate:.2%}",
            "下场动作": "核查退款集中货号，未确认原因前不作为主推款",
        })
    room_talks = talks[talks["live_room_id"].astype(str) == room].copy() if not talks.empty else pd.DataFrame()
    if not room_talks.empty:
        declining = room_talks.groupby("style_code", as_index=False).agg(
            讲解分钟=("讲解分钟", "sum"), 平台支付=("paid_amount", "sum"),
            在线人数变化=("viewer_change", "sum"), 成交件数=("sold_units", "sum"),
        )
        declining["分钟产出"] = declining["平台支付"].div(declining["讲解分钟"].replace(0, pd.NA))
        for _, action_row in declining[(declining["讲解分钟"] >= 5) & (declining["在线人数变化"] < 0)].head(3).iterrows():
            room_actions.append({
                "优先级": "中", "本场发现": f"{action_row['style_code']}讲解{action_row['讲解分钟']:.1f}分钟，在线人数净变化{action_row['在线人数变化']:+.0f}",
                "下场动作": "缩短讲解或后移至非流量高峰，并重新测试话术",
            })
    if not room_products.empty:
        best_product = room_products.sort_values("paid_amount", ascending=False).iloc[0]
        if float(best_product.get("paid_amount", 0) or 0) > 0:
            room_actions.append({
                "优先级": "机会", "本场发现": f"{best_product.get('style_code') or '未识别货号'}为本场支付最高商品，平台支付¥{float(best_product['paid_amount']):,.0f}",
                "下场动作": "保留该商品并测试提前上场，结合多场稳定性决定是否进入固定排品",
            })
    if room_actions:
        show_table(pd.DataFrame(room_actions))
    else:
        st.info("本场未触发高优先级规则，可进入商品明细继续检查。")
    detail_tabs = st.tabs(["商品平台表现", "商品讲解区间", "渠道流量", "全部直播指标"])
    with detail_tabs[0]:
        show_table(room_products[["style_code", "product_name", "click_users", "sold_units", "paid_amount", "pre_ship_refund_amount", "post_ship_refund_amount"]].sort_values("paid_amount", ascending=False))
    with detail_tabs[1]:
        show_table(talks[talks["live_room_id"].astype(str) == room].sort_values("talk_start_epoch") if not talks.empty else pd.DataFrame())
    with detail_tabs[2]:
        room_channels = channels[channels["live_room_id"].astype(str) == room].copy() if not channels.empty else pd.DataFrame()
        if room_channels.empty:
            st.info("本场暂无渠道流量数据。")
        else:
            detail_channels = room_channels[room_channels["channel_name"] != "整体"].copy()
            detail_channels["观看人数占比"] = detail_channels["watch_users"].div(
                detail_channels["watch_users"].sum() or pd.NA
            )
            detail_channels["支付金额占比"] = detail_channels["paid_amount"].div(
                detail_channels["paid_amount"].sum() or pd.NA
            )
            detail_channels["千次观看支付"] = detail_channels["paid_amount"].div(
                detail_channels["watch_users"].replace(0, pd.NA)
            ) * 1000
            detail_channels["投放消耗观察值"] = detail_channels[["shop_bound_spend", "shop_promoted_spend"]].max(axis=1)
            detail_channels["ROI"] = detail_channels["paid_amount"].div(
                detail_channels["投放消耗观察值"].replace(0, pd.NA)
            )
            show_table(detail_channels.sort_values("paid_amount", ascending=False))
            st.caption("渠道投放指标用于经营观察；投放归因口径与平台支付口径可能不同，不作为财务ROI。")
    with detail_tabs[3]:
        show_table(metrics[metrics["live_room_id"].astype(str) == room][["module", "metric_name", "raw_value", "comparison_display"]])

with compare_tab:
    room_compare = sessions.groupby(["shop_name", "anchor_name"], as_index=False).agg(场次=("live_room_id", "nunique"), 直播小时=("duration_seconds", lambda x: x.sum()/3600))
    performance = products.groupby(["shop_name", "anchor_name"], as_index=False).agg(商品点击=("click_users", "sum"), 成交件数=("sold_units", "sum"), 平台支付=("paid_amount", "sum"))
    room_compare = room_compare.merge(performance, on=["shop_name", "anchor_name"], how="left")
    compare_kpis = session_kpis.copy()
    compare_kpis["观看时长加权"] = compare_kpis["人均观看时长"] * compare_kpis["直播间观看人数"]
    compare_kpis = compare_kpis.groupby(["shop_name", "anchor_name"], as_index=False).agg(
        曝光人数=("直播间曝光人数", "sum"), 观看人数=("直播间观看人数", "sum"),
        平均在线=("平均在线人数", "mean"), 最高在线=("最高在线人数", "max"),
        自然观看=("自然流量观看人数", "sum"), 付费观看=("付费流量观看人数", "sum"),
        成交人数=("直播间成交人数", "sum"), 新增粉丝=("新增粉丝数", "sum"),
        观看时长加权=("观看时长加权", "sum"), 投放消耗观察值=("投放消耗观察值", "sum"),
    )
    room_compare = room_compare.merge(compare_kpis, on=["shop_name", "anchor_name"], how="left")
    room_compare["曝光进入率"] = room_compare["观看人数"].div(room_compare["曝光人数"].replace(0, pd.NA))
    room_compare["人均观看时长"] = room_compare["观看时长加权"].div(room_compare["观看人数"].replace(0, pd.NA))
    room_compare["商品点击率"] = room_compare["商品点击"].div(room_compare["观看人数"].replace(0, pd.NA))
    room_compare["观看成交率"] = room_compare["成交人数"].div(room_compare["观看人数"].replace(0, pd.NA))
    room_compare["关注转化率"] = room_compare["新增粉丝"].div(room_compare["观看人数"].replace(0, pd.NA))
    room_compare["自然流量占比"] = room_compare["自然观看"].div(
        (room_compare["自然观看"] + room_compare["付费观看"]).replace(0, pd.NA)
    )
    room_compare["场均支付"] = room_compare["平台支付"].div(room_compare["场次"].replace(0, pd.NA))
    room_compare["每小时支付"] = room_compare["平台支付"].div(room_compare["直播小时"].replace(0, pd.NA))
    room_compare["点击成交率"] = room_compare["成交件数"].div(room_compare["商品点击"].replace(0, pd.NA))
    room_compare["每小时新增粉丝"] = room_compare["新增粉丝"].div(room_compare["直播小时"].replace(0, pd.NA))
    room_compare["ROI"] = room_compare["平台支付"].div(room_compare["投放消耗观察值"].replace(0, pd.NA))
    room_compare_display = room_compare.drop(
        columns=["观看时长加权", "投放消耗观察值"], errors="ignore"
    )
    show_table(room_compare_display.sort_values("平台支付", ascending=False))
    st.plotly_chart(px.bar(room_compare, x="shop_name", y="每小时支付", color="anchor_name", title="直播间每小时产出对比", labels={"shop_name": "直播间／店铺", "anchor_name": "主播"}), width="stretch")
    st.markdown("#### 同一商品的主播适配")
    compare_style_options = style_summary.sort_values("平台支付", ascending=False)["style_code"].astype(str).tolist()
    if compare_style_options:
        compare_style = st.selectbox("选择货号进行公平对比", compare_style_options, key="anchor_compare_style")
        compare_product = products[products["style_code"].astype(str) == compare_style].copy()
        same_product_compare = compare_product.groupby(["anchor_name", "shop_name"], as_index=False).agg(
            上播场次=("live_room_id", "nunique"), 商品点击=("click_users", "sum"),
            成交件数=("sold_units", "sum"), 平台支付=("paid_amount", "sum"),
        )
        same_product_compare["点击成交率"] = same_product_compare["成交件数"].div(
            same_product_compare["商品点击"].replace(0, pd.NA)
        )
        same_product_compare["场均支付"] = same_product_compare["平台支付"].div(
            same_product_compare["上播场次"].replace(0, pd.NA)
        )
        same_product_compare["主播占该货号成交"] = same_product_compare["成交件数"].div(
            same_product_compare["成交件数"].sum() or pd.NA
        )
        show_table(same_product_compare.sort_values(["点击成交率", "场均支付"], ascending=False))
        st.caption("同货号对比优先看点击成交率和场均支付，避免仅用总销售额评价主播。")

@st.fragment
def render_session_trend(trend_frame: pd.DataFrame) -> None:
    """指标切换只重绘趋势组件，不触发整页直播分析重新计算。"""
    metric_name = st.selectbox("趋势指标", [
        "平台支付", "直播间观看人数", "曝光进入率", "人均观看时长", "平均在线人数", "互动率",
        "商品点击率", "点击成交率", "新增粉丝数", "关注转化率", "自然流量占比",
        "首购率", "粉丝成交人数占比", "千次观看用户支付金额",
        "投放消耗观察值", "ROI",
    ], key="session_trend_metric")
    st.plotly_chart(px.line(
        trend_frame.sort_values("start_time"), x="start_time", y=metric_name, color="anchor_name",
        markers=True, hover_data=["shop_name", "live_room_id"],
        labels={"start_time": "开播时间", "anchor_name": "主播", "shop_name": "直播间／店铺", "live_room_id": "直播场次ID"},
    ), width="stretch", key="session_trend_chart")


with decision_tab:
    st.markdown("#### 场次经营趋势")
    trend_frame = session_kpis.rename(columns={
        "商品支付金额": "平台支付", "商品点击人数": "商品点击", "商品成交件数": "成交件数",
    }).copy()
    render_session_trend(trend_frame)

with product_tab:
    product_filter_cols = st.columns([2, 1])
    with product_filter_cols[0]:
        search = st.text_input("搜索商品名称或货号")
    with product_filter_cols[1]:
        product_min_sessions = st.number_input(
            "最少上播场次",
            min_value=1,
            max_value=max(1, int(safe_number(style_summary["上播场次"].max(), 1))),
            value=1,
            step=1,
            key="product_decision_min_sessions",
        )
    table = style_summary.copy()
    table = table[table["上播场次"] >= product_min_sessions]
    if search:
        key = search.strip().upper()
        table = table[table["style_code"].astype(str).str.upper().str.contains(key, regex=False) | table["商品名称"].astype(str).str.upper().str.contains(key, regex=False)]
    table["诊断"] = "继续观察"
    table.loc[(table["上播场次"] >= 3) & (table["点击成交率"] >= .06), "诊断"] = "稳定转化"
    table.loc[(table["累计点击"] >= 100) & (table["点击成交率"] < .03), "诊断"] = "高点击低成交"
    table.loc[table["退款率"] >= .35, "诊断"] = "退款风险"
    show_table(table.sort_values("平台支付", ascending=False))

with selection_tab:
    st.subheader("开播前段商品销售排行")
    rank_cols = st.columns([1.2, 1.2, 1.4, 1.5])
    with rank_cols[0]:
        rank_window = st.segmented_control(
            "统计范围", [15, 30, 60, 90], default=60,
            format_func=lambda value: f"前{value}分钟", key="opening_sales_window",
        )
    with rank_cols[1]:
        rank_by = st.selectbox(
            "排序指标", ["前段成交金额", "前段成交件数", "成交场次率", "支付/讲解分钟", "在线上升场次占比"]
        )
    with rank_cols[2]:
        minimum_room_count = st.slider("最少出现直播场次", 1, 10, 1, key="opening_min_rooms")
    with rank_cols[3]:
        launch_day_mode = st.selectbox(
            "上新当天口径", ["全部销售", "排除上新当天", "只看上新当天"],
            key="opening_launch_day_mode",
        )

    opening_talks = talks[talks["开播后分钟"] <= rank_window].copy() if not talks.empty else pd.DataFrame()
    if opening_talks.empty:
        st.warning("当前范围没有商品讲解区间数据，上传新版直播工作簿后才能生成排行。")
    else:
        opening_room = opening_talks.groupby(["style_code", "live_room_id"], as_index=False).agg(
            直播日期=("session_start_utc", lambda values: values.iloc[0].date() if len(values) else None),
            讲解次数=("product_id", "size"), 讲解分钟=("讲解分钟", "sum"),
            前段成交金额=("paid_amount", "sum"), 前段成交件数=("sold_units", "sum"),
            在线净变化=("viewer_change", "sum"), 分均在线人数=("avg_online_users", "mean"),
        )
        opening_room["有成交"] = opening_room["前段成交件数"] > 0
        opening_room["在线上升"] = opening_room["在线净变化"] > 0
        launch_lookup = style_summary[["style_code", "上新日期"]].drop_duplicates("style_code") if "上新日期" in style_summary else pd.DataFrame(columns=["style_code", "上新日期"])
        opening_room = opening_room.merge(launch_lookup, on="style_code", how="left")
        opening_room["上新日期"] = pd.to_datetime(opening_room["上新日期"], errors="coerce").dt.date
        opening_room["上新当天"] = opening_room["上新日期"].notna() & (opening_room["直播日期"] == opening_room["上新日期"])
        opening_room["日期异常"] = opening_room["上新日期"].notna() & (opening_room["直播日期"] < opening_room["上新日期"])

        launch_effect = opening_room.groupby("style_code", as_index=False).agg(
            原始前段成交金额=("前段成交金额", "sum"),
            上新当天成交金额=("前段成交金额", lambda values: values[opening_room.loc[values.index, "上新当天"]].sum()),
            原始前段成交件数=("前段成交件数", "sum"),
            上新当天成交件数=("前段成交件数", lambda values: values[opening_room.loc[values.index, "上新当天"]].sum()),
        )
        launch_effect["排除后成交金额"] = launch_effect["原始前段成交金额"] - launch_effect["上新当天成交金额"]
        launch_effect["排除后成交件数"] = launch_effect["原始前段成交件数"] - launch_effect["上新当天成交件数"]
        launch_effect["上新当天金额占比"] = launch_effect["上新当天成交金额"].div(launch_effect["原始前段成交金额"].replace(0, pd.NA))

        analysis_room = opening_room[~opening_room["日期异常"]].copy()
        if launch_day_mode == "排除上新当天":
            analysis_room = analysis_room[~analysis_room["上新当天"]]
        elif launch_day_mode == "只看上新当天":
            analysis_room = analysis_room[analysis_room["上新当天"]]

        opening_rank = analysis_room.groupby("style_code", as_index=False).agg(
            前段出现场次=("live_room_id", "nunique"), 前段成交场次=("有成交", "sum"),
            前段讲解次数=("讲解次数", "sum"), 前段讲解分钟=("讲解分钟", "sum"),
            前段成交金额=("前段成交金额", "sum"), 前段成交件数=("前段成交件数", "sum"),
            在线净变化=("在线净变化", "sum"), 平均在线变化=("在线净变化", "mean"),
            在线上升场次占比=("在线上升", "mean"), 分均在线人数=("分均在线人数", "mean"),
        )
        opening_rank["成交场次率"] = opening_rank["前段成交场次"].div(opening_rank["前段出现场次"].replace(0, pd.NA))
        opening_rank["支付/讲解分钟"] = opening_rank["前段成交金额"].div(opening_rank["前段讲解分钟"].replace(0, pd.NA))
        opening_rank["成交件数/讲解分钟"] = opening_rank["前段成交件数"].div(opening_rank["前段讲解分钟"].replace(0, pd.NA))
        context_column_names = [
            "style_code", "商品名称", "累计点击", "平台退款", "退款率",
            "ship_amount", "return_amount", "net_amount", "品牌", "年份", "品类", "上新日期",
        ]
        context_columns = style_summary[[column for column in context_column_names if column in style_summary.columns]]
        opening_rank = opening_rank.merge(context_columns, on="style_code", how="left").merge(
            launch_effect, on="style_code", how="left"
        )
        opening_output_median = opening_rank.loc[
            opening_rank["前段讲解分钟"] > 0, "支付/讲解分钟"
        ].median()
        opening_output_median = float(opening_output_median) if pd.notna(opening_output_median) else 0

        def opening_recommendation(row):
            room_count = int(safe_number(row.get("前段出现场次")))
            conversion_rooms = safe_number(row.get("成交场次率"))
            online_up = safe_number(row.get("在线上升场次占比"))
            output = safe_number(row.get("支付/讲解分钟"))
            refund = safe_number(row.get("退款率"))
            confidence = "高可信" if room_count >= 5 else "中可信" if room_count >= 3 else "待验证"
            if refund >= .35:
                return pd.Series(["暂缓主推", "退款风险较高，先核查原因", confidence])
            if conversion_rooms >= .6 and online_up >= .5:
                return pd.Series([f"前{rank_window}分钟主力", "多场成交且在线承接稳定", confidence])
            if opening_output_median and output >= opening_output_median * 1.3:
                return pd.Series(["极速流承接", "讲解分钟产出高，适合流量到来时快速承接", confidence])
            if conversion_rooms >= .4:
                return pd.Series(["常规测试", "已有重复成交，继续积累样本", confidence])
            return pd.Series(["观察／后移", "前段成交稳定性不足，暂不占用核心时段", confidence])

        if opening_rank.empty:
            opening_rank[["推荐场景", "建议依据", "可信度"]] = pd.DataFrame(
                columns=["推荐场景", "建议依据", "可信度"], index=opening_rank.index
            )
        else:
            opening_rank[["推荐场景", "建议依据", "可信度"]] = opening_rank.apply(opening_recommendation, axis=1)
        dimension_cols = st.columns(3)
        selected_brands = dimension_cols[0].multiselect(
            "品牌", sorted(opening_rank["品牌"].dropna().astype(str).unique()), key="opening_brands"
        ) if "品牌" in opening_rank else []
        selected_categories = dimension_cols[1].multiselect(
            "品类", sorted(opening_rank["品类"].dropna().astype(str).unique()), key="opening_categories"
        ) if "品类" in opening_rank else []
        selected_years = dimension_cols[2].multiselect(
            "年份", sorted(opening_rank["年份"].dropna().astype(str).unique()), key="opening_years"
        ) if "年份" in opening_rank else []
        if selected_brands:
            opening_rank = opening_rank[opening_rank["品牌"].astype(str).isin(selected_brands)]
        if selected_categories:
            opening_rank = opening_rank[opening_rank["品类"].astype(str).isin(selected_categories)]
        if selected_years:
            opening_rank = opening_rank[opening_rank["年份"].astype(str).isin(selected_years)]
        opening_rank = opening_rank[opening_rank["前段出现场次"] >= minimum_room_count]
        opening_rank = opening_rank.sort_values(rank_by, ascending=False).reset_index(drop=True)
        opening_rank.insert(0, "排名", opening_rank.index + 1)

        total_opening_pay = opening_rank["前段成交金额"].sum()
        total_opening_units = opening_rank["前段成交件数"].sum()
        top_ten_share = opening_rank.head(10)["前段成交金额"].sum() / total_opening_pay if total_opening_pay else 0
        total_talk_minutes = opening_rank["前段讲解分钟"].sum()
        summary_cards = st.columns(5)
        summary_cards[0].metric(f"前{rank_window}分钟成交金额", f"¥{total_opening_pay:,.0f}")
        summary_cards[1].metric("成交件数", f"{total_opening_units:,.0f}")
        summary_cards[2].metric("有成交商品", f"{(opening_rank['前段成交件数'] > 0).sum():,}")
        summary_cards[3].metric("前10商品金额占比", f"{top_ten_share:.2%}")
        summary_cards[4].metric("平均分钟产出", f"¥{total_opening_pay / total_talk_minutes:,.0f}" if total_talk_minutes else "—")

        unique_opening_styles = opening_room["style_code"].nunique()
        dated_opening_styles = opening_room.loc[opening_room["上新日期"].notna(), "style_code"].nunique()
        launch_day_pay = opening_room.loc[opening_room["上新当天"], "前段成交金额"].sum()
        launch_cards = st.columns(3)
        launch_cards[0].metric("上新日期识别覆盖", f"{dated_opening_styles}/{unique_opening_styles}款")
        launch_cards[1].metric("上新当天前段成交", f"¥{launch_day_pay:,.0f}")
        launch_cards[2].metric("排除后前段成交", f"¥{opening_room['前段成交金额'].sum() - launch_day_pay:,.0f}")

        st.caption(
            f"“前{rank_window}分钟”按讲解开始时间归入；跨越边界的讲解区间整段计入。"
            f"当前上新口径：{launch_day_mode}。直播日期等于商品上新日期时判定为上新当天；"
            "上新日期晚于直播日期的数据标记异常并排除。累计点击、平台退款为整场指标；"
            "范围发货、退货和实销为所选日期范围履约数据。"
        )
        display_rank = opening_rank.rename(columns={
            "style_code": "货号", "累计点击": "整场累计点击", "平台退款": "整场平台退款",
            "退款率": "整场退款率", "ship_amount": "范围发货", "return_amount": "范围退货",
            "net_amount": "范围实销",
        })
        show_table(display_rank)
        st.download_button(
            f"下载前{rank_window}分钟商品销售排行",
            display_rank.to_csv(index=False).encode("utf-8-sig"),
            f"直播前{rank_window}分钟商品销售排行_{start_date}_{end_date}.csv",
            "text/csv", key="download_opening_sales_rank",
        )

        if opening_rank.empty:
            st.info("当前筛选条件下没有符合要求的商品。")
        else:
            evidence_style = st.selectbox(
                "查看货号逐场证据", opening_rank["style_code"].astype(str).tolist(),
                format_func=lambda code: f"{code}｜{opening_rank.loc[opening_rank['style_code'].astype(str) == code, '商品名称'].iloc[0]}",
                key="opening_evidence_style",
            )
            evidence = analysis_room[analysis_room["style_code"].astype(str) == evidence_style].merge(
                sessions[["live_room_id", "shop_name", "anchor_name"]], on="live_room_id", how="left"
            )
            show_table(evidence.sort_values("直播日期", ascending=False))

with selection_tab:
    st.markdown("#### 商品机会类型")
    opportunity_help = {
        "开播阶段": "查看开播后前15／30／60／90分钟内表现较好的商品，辅助安排开场和前段排品。",
        "讲解高效": "主要参考讲解时长、支付金额和成交件数，识别单位讲解时间产出较高的商品。",
        "在线提升": "观察商品讲解期间在线人数变化，识别更容易稳住或提升直播间在线人数的商品。",
        "高转化": "主要参考商品点击人数与成交件数，识别点击后成交效率较高的商品。",
        "高引流": "主要参考商品累计点击人数，识别更容易吸引用户点击和了解的商品。",
        "稳定复销": "判断商品是否在多场、跨周持续成交，并结合成交场次率和最近3场表现，避免把单场爆发误判为稳定款。",
        "高点击低成交": "商品获得较多点击但成交率偏低，通常需要检查价格、利益点、尺码说明或讲解方式。",
        "退款风险": "按平台退款率排序，辅助识别成交后退款风险较高、需要核查商品或讲解问题的款式。",
    }
    opportunity_types = list(opportunity_help)
    if st.session_state.get("live_opportunity_type") not in opportunity_types:
        st.session_state["live_opportunity_type"] = "开播阶段"
    opportunity_buttons = st.columns(len(opportunity_types))
    for button_column, opportunity_name in zip(opportunity_buttons, opportunity_types):
        with button_column:
            if st.button(
                f"{opportunity_name}  ⓘ",
                key=f"opportunity_{opportunity_name}",
                help=opportunity_help[opportunity_name],
                type="primary" if st.session_state["live_opportunity_type"] == opportunity_name else "secondary",
                width="stretch",
            ):
                st.session_state["live_opportunity_type"] = opportunity_name
                st.rerun()
    opportunity = st.session_state["live_opportunity_type"]
    min_sessions = st.slider("最少上播场次", 1, 10, 3)
    min_clicks = st.slider("最少累计点击", 0, 1000, 50, 10)
    candidates = style_summary.copy()
    if opportunity == "开播阶段":
        opening_window = st.segmented_control("开播后范围", [15, 30, 60, 90], default=60, format_func=lambda value: f"前{value}分钟")
        opening_talks = talks[talks["开播后分钟"] <= opening_window].copy() if not talks.empty else pd.DataFrame()
        if opening_talks.empty:
            candidates = candidates.iloc[0:0]
        else:
            opening = opening_talks.groupby("style_code", as_index=False).agg(
                开播阶段场次=("live_room_id", "nunique"), 开播阶段讲解次数=("product_id", "size"),
                开播阶段讲解分钟=("讲解分钟", "sum"), 开播阶段支付=("paid_amount", "sum"),
                开播阶段成交件数=("sold_units", "sum"), 开播阶段在线净变化=("viewer_change", "sum"),
                开播阶段平均在线变化=("viewer_change", "mean"), 开播阶段在线上升占比=("在线上升", "mean"),
            )
            opening["开播阶段分钟产出"] = opening["开播阶段支付"].div(opening["开播阶段讲解分钟"].replace(0, pd.NA))
            candidates = candidates.merge(opening, on="style_code", how="inner")
            candidates = candidates[candidates["开播阶段场次"] >= min_sessions].sort_values(
                ["开播阶段分钟产出", "开播阶段在线上升占比"], ascending=False
            )
        st.caption(f"只统计商品讲解开始时间位于开播后前{opening_window}分钟的区间。")
    else:
        if opportunity != "稳定复销":
            candidates = candidates[(candidates["上播场次"] >= min_sessions) & (candidates["累计点击"] >= min_clicks)]
    if opportunity == "讲解高效": candidates = candidates.sort_values(["支付/讲解分钟", "成交件数/讲解分钟"], ascending=False)
    elif opportunity == "在线提升": candidates = candidates.sort_values(["在线上升场次占比", "平均在线变化"], ascending=False)
    elif opportunity == "高转化": candidates = candidates.sort_values("点击成交率", ascending=False)
    elif opportunity == "高引流": candidates = candidates.sort_values("累计点击", ascending=False)
    elif opportunity == "稳定复销":
        stability_order = {
            "高置信稳定": 0, "稳定复销": 1, "有复销潜力": 2,
            "单场爆发": 3, "近期转弱": 4, "待观察": 5, "样本不足": 6,
        }
        candidates["复销排序"] = candidates["复销分级"].map(stability_order).fillna(9)
        candidates = candidates.sort_values(
            ["复销排序", "复销成交场次率", "复销上播场次"], ascending=[True, False, False]
        ).drop(columns="复销排序")
        level_counts = candidates["复销分级"].value_counts()
        stability_levels = ["高置信稳定", "稳定复销", "有复销潜力", "单场爆发", "近期转弱", "待观察", "样本不足"]
        level_options = [f"{label} · {int(level_counts.get(label, 0))}" for label in stability_levels]
        selected_level_option = st.segmented_control(
            "复销分类（点击查看具体款式）", level_options,
            default=f"稳定复销 · {int(level_counts.get('稳定复销', 0))}",
            key="stability_level_filter",
        )
        selected_stability_level = selected_level_option.rsplit(" · ", 1)[0]
        candidates = candidates[candidates["复销分级"] == selected_stability_level].copy()
        st.caption(
            "稳定复销：至少3场、至少2场成交、成交场次率≥50%，且最近3场仍有成交。"
            "高置信稳定：至少5场、至少3场成交、成交场次率≥60%、跨2周，且最近3场至少2场成交。"
            "退款率仅作为风险提示，不再否决复销分级。"
        )
    elif opportunity == "高点击低成交": candidates = candidates[candidates["点击成交率"] < .03].sort_values("累计点击", ascending=False)
    elif opportunity == "退款风险": candidates = candidates.sort_values("退款率", ascending=False)
    if opportunity in ["讲解高效", "在线提升", "开播阶段"] and talks.empty:
        st.warning("当前范围没有新版商品讲解区间数据，上传新版直播工作簿后才能计算。")
    show_table(candidates)
    if opportunity == "稳定复销":
        if candidates.empty:
            st.info(f"当前没有“{selected_stability_level}”商品。")
        else:
            repeat_style = st.selectbox(
                "选择货号查看逐场证据",
                candidates["style_code"].astype(str).tolist(),
                format_func=lambda code: (
                    f"{code}｜{candidates.loc[candidates['style_code'].astype(str) == code, '商品名称'].iloc[0]}"
                ),
                key="repeat_sales_evidence_style",
            )
            repeat_evidence = session_style[session_style["style_code"].astype(str) == repeat_style].merge(
                sessions[["live_room_id", "shop_name", "anchor_name"]], on="live_room_id", how="left"
            )
            repeat_evidence["该场成交占比"] = repeat_evidence["场次成交件数"].div(
                repeat_evidence["场次成交件数"].sum() or pd.NA
            )
            show_table(repeat_evidence.sort_values("直播日期", ascending=False))

with product_tab:
    options = style_summary.sort_values("平台支付", ascending=False)["style_code"].astype(str).tolist()
    selected_style = st.selectbox("选择货号", options)
    with st.expander(f"{selected_style} 库存明细"):
        render_inventory_detail(selected_style, f"live_inventory_{selected_style}")
    item = products[products["style_code"].astype(str) == selected_style].copy()
    summary_row = style_summary[style_summary["style_code"].astype(str) == selected_style].iloc[0]
    cards = st.columns(5)
    cards[0].metric("历史上播场次", f"{summary_row['上播场次']:.0f}")
    cards[1].metric("成交场次率", f"{summary_row['成交场次率']:.2%}")
    cards[2].metric("点击成交率", f"{summary_row['点击成交率']:.2%}")
    cards[3].metric("平台支付", f"¥{summary_row['平台支付']:,.0f}")
    cards[4].metric("范围实销", f"¥{summary_row['net_amount']:,.0f}")
    product_level = safe_text(summary_row.get("复销分级"), "样本不足")
    product_refund = safe_number(summary_row.get("退款率"))
    product_conversion = safe_number(summary_row.get("点击成交率"))
    if product_refund >= .35:
        product_decision = "暂缓主推：退款风险较高，先核查退款原因和商品反馈。"
    elif product_level in {"高置信稳定", "稳定复销"}:
        product_decision = f"建议保留在常规排品：当前判定为{product_level}，可继续测试提前上场。"
    elif product_conversion < .03 and safe_number(summary_row.get("累计点击")) >= 100:
        product_decision = "建议优化后复测：点击充足但成交偏弱，重点检查价格、利益点和讲解表达。"
    elif product_level == "单场爆发":
        product_decision = "暂不判断为稳定款：成交依赖少数场次，需要继续验证复销能力。"
    else:
        product_decision = "继续积累样本：目前不足以形成明确主推或淘汰结论。"
    st.info(f"商品决策：{product_decision}")
    anchor_item = item.groupby("anchor_name", as_index=False).agg(场次=("live_room_id", "nunique"), 点击=("click_users", "sum"), 成交件数=("sold_units", "sum"), 平台支付=("paid_amount", "sum"))
    anchor_item["点击成交率"] = anchor_item["成交件数"].div(anchor_item["点击"].replace(0, pd.NA))
    anchor_item["该货号占主播成交"] = anchor_item.apply(lambda r: r["成交件数"] / products.loc[products["anchor_name"] == r["anchor_name"], "sold_units"].sum() if products.loc[products["anchor_name"] == r["anchor_name"], "sold_units"].sum() else pd.NA, axis=1)
    anchor_item["主播占该货号成交"] = anchor_item["成交件数"].div(anchor_item["成交件数"].sum() or pd.NA)
    st.subheader("主播适配与占比")
    show_table(anchor_item.sort_values("成交件数", ascending=False))
    st.subheader("讲解效率与在线变化")
    item_talks = talks[talks["style_code"].astype(str) == selected_style].copy() if not talks.empty else pd.DataFrame()
    if item_talks.empty:
        st.info("该货号暂无新版商品讲解区间数据。")
    else:
        efficiency = st.columns(5)
        talk_minutes = item_talks["讲解分钟"].sum()
        efficiency[0].metric("累计讲解时长", f"{talk_minutes:,.1f}分钟")
        efficiency[1].metric("场均讲解时长", f"{talk_minutes / item_talks['live_room_id'].nunique():,.1f}分钟")
        efficiency[2].metric("支付/讲解分钟", f"¥{item_talks['paid_amount'].sum() / talk_minutes:,.0f}" if talk_minutes else "—")
        efficiency[3].metric("平均在线变化", f"{item_talks['viewer_change'].mean():+,.1f}")
        efficiency[4].metric("在线上升区间占比", f"{item_talks['在线上升'].mean():.2%}")
        show_table(
            item_talks[["live_room_id", "开播后分钟", "讲解分钟", "paid_amount", "sold_units", "viewer_change", "avg_online_users"]]
            .sort_values(["live_room_id", "开播后分钟"])
        )
    st.subheader("逐场历史")
    show_table(item[["直播日期", "shop_name", "anchor_name", "talk_count", "click_users", "sold_units", "paid_amount", "pre_ship_refund_amount", "post_ship_refund_amount"]].sort_values("直播日期", ascending=False))

def _export_frame(value) -> pd.DataFrame:
    if value is None:
        return pd.DataFrame()
    if isinstance(value, pd.Series):
        frame = value.to_frame().T
    elif isinstance(value, pd.DataFrame):
        frame = value.copy()
    elif isinstance(value, list):
        frame = pd.DataFrame(value)
    elif isinstance(value, dict):
        frame = pd.DataFrame([value])
    else:
        frame = pd.DataFrame({"内容": [value]})
    if frame.empty:
        return frame
    frame = localize_table(frame)
    for column in frame.columns:
        if isinstance(frame[column].dtype, pd.DatetimeTZDtype):
            frame[column] = frame[column].dt.tz_localize(None)
        elif pd.api.types.is_datetime64_any_dtype(frame[column]):
            frame[column] = pd.to_datetime(frame[column], errors="coerce")
        elif frame[column].dtype == "object":
            frame[column] = frame[column].map(
                lambda value: json.dumps(
                    sorted(value) if isinstance(value, set) else value,
                    ensure_ascii=False,
                )
                if isinstance(value, (dict, list, tuple, set)) else value
            )
    return frame


def _current_export_sheets() -> dict[str, pd.DataFrame]:
    sheets = {
        "筛选条件": pd.DataFrame([{
            "当前栏目": active_section,
            "分析周期": quick,
            "开始日期": start_date,
            "结束日期": end_date,
            "店铺": "、".join(selected_shops),
            "主播": "、".join(selected_anchors),
            "导出口径": "平台指标按直播发生时间；发货、退货和实销按仓库发生时间",
        }]),
    }
    scope = globals()
    if active_section == "决策总览":
        sheets.update({
            "经营链路": pd.DataFrame([{
                "曝光进入率": overall_entry_rate, "人均观看时长（秒）": weighted_watch_seconds,
                "商品点击率": overall_click_rate, "点击成交率": overall_click_conversion,
                "关注转化率": overall_follow_rate, "ROI": overall_spend_output,
            }]),
            "优先处理清单": decisions,
            "下一场建议": scope.get("next_display"),
            "每日经营趋势": daily_platform,
            "逐场经营链路": session_kpis,
            "全部商品判断": style_summary,
            "判断证据": scope.get("evidence_rows"),
        })
    elif active_section == "开播选品":
        sheets.update({
            "前段商品排行": scope.get("display_rank"),
            "当前机会类型": pd.DataFrame([{
                "机会类型": opportunity, "最少上播场次": min_sessions,
                "最少累计点击": min_clicks,
            }]),
            "机会商品": candidates,
            "讲解汇总": talk_summary,
            "前段逐场数据": scope.get("analysis_room"),
            "前段判断证据": scope.get("evidence"),
            "复销判断证据": scope.get("repeat_evidence"),
            "全部商品判断": style_summary,
        })
    elif active_section == "单场复盘":
        selected_room_talks = talks[talks["live_room_id"].astype(str) == room].copy() if not talks.empty else pd.DataFrame()
        selected_room_metrics = metrics[metrics["live_room_id"].astype(str) == room].copy() if not metrics.empty else pd.DataFrame()
        sheets.update({
            "场次信息": labels[labels["live_room_id"].astype(str) == room],
            "本场经营链路": room_kpi,
            "本场诊断建议": room_actions,
            "商品平台表现": room_products,
            "商品讲解区间": selected_room_talks,
            "渠道流量": scope.get("detail_channels", scope.get("room_channels")),
            "全部直播指标": selected_room_metrics,
        })
    elif active_section == "商品决策":
        sheets.update({
            "当前商品结论": pd.DataFrame([{"货号": selected_style, "商品决策": product_decision}]),
            "当前商品汇总": summary_row,
            "筛选后商品汇总": table,
            "主播适配": anchor_item,
            "讲解区间": item_talks,
            "逐场历史": item,
        })
    elif active_section == "主播对比":
        sheets.update({
            "主播经营对比": room_compare_display,
            "同商品主播对比": scope.get("same_product_compare"),
            "场次趋势底表": trend_frame,
            "全部商品汇总": style_summary,
        })
    prepared = {}
    for name, frame in sheets.items():
        prepared_frame = _export_frame(frame)
        if not prepared_frame.empty:
            prepared[name] = prepared_frame
    return prepared


def _build_live_export(sheets: dict[str, pd.DataFrame]) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            safe_name = sheet_name[:31]
            frame.to_excel(writer, index=False, sheet_name=safe_name)
            worksheet = writer.book[safe_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="17324D")
            for column_cells in worksheet.columns:
                width = min(45, max(10, max(len(str(cell.value or "")) for cell in column_cells[:200]) + 2))
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
            for column_index, column_name in enumerate(frame.columns, start=1):
                if "率" in str(column_name) or "占比" in str(column_name):
                    for cell in worksheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                        for value_cell in cell:
                            if isinstance(value_cell.value, (int, float)):
                                value_cell.number_format = "0.00%"
                elif str(column_name) == "ROI":
                    for cell in worksheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                        for value_cell in cell:
                            if isinstance(value_cell.value, (int, float)):
                                value_cell.number_format = "0.00"
    return output.getvalue()


@st.fragment
def render_current_page_export(sheets: dict[str, pd.DataFrame]) -> None:
    export_key = f"_live_export_{active_section}_{start_date}_{end_date}"
    if st.button("生成当前页面完整报表", key="build_current_live_export", type="primary"):
        with st.spinner("正在整理当前页面全部信息……"):
            st.session_state[export_key] = _build_live_export(sheets)
    export_data = st.session_state.get(export_key)
    if export_data:
        st.download_button(
            "下载当前页面全部信息",
            export_data,
            file_name=f"直播经营分析_{active_section}_{start_date}_{end_date}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="download_current_live_export",
        )


st.markdown("---")
render_current_page_export(_current_export_sheets())

for hidden_holder in hidden_section_holders:
    hidden_holder.empty()
